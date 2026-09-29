#!/usr/bin/env python3
"""Build the Beto fuel dataset from ANP's official weekly summary."""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import urllib.request
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable
from urllib.parse import urljoin

from openpyxl import load_workbook

SOURCE_INDEX = "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/precos/levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas"
SOURCE_NAME = "ANP — Levantamento de Preços de Combustíveis (resumo semanal)"
SCHEMA_VERSION = 1
MIN_RECORDS = 30
MIN_PRICE = 0.05
MAX_PRICE = 100.0
MAX_PRICE_PER_13KG = 300.0
ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"

FUEL_ALIASES = {
    "GASOLINA COMUM": ("GASOLINE_C", "R$/L"),
    "GASOLINA C": ("GASOLINE_C", "R$/L"),
    "GASOLINA COMUM": ("GASOLINE_C", "R$/L"),
    "GASOLINA ADITIVADA": ("GASOLINE_ADDITIVATED", "R$/L"),
    "GASOLINA C ADITIVADA": ("GASOLINE_ADDITIVATED", "R$/L"),
    "GASOLINA PREMIUM": ("GASOLINE_PREMIUM", "R$/L"),
    "ETANOL HIDRATADO": ("ETHANOL_HYDRATED", "R$/L"),
    "ETANOL": ("ETHANOL_HYDRATED", "R$/L"),
    "GNV": ("GNV", "R$/m³"),
    "GÁS NATURAL VEICULAR": ("GNV", "R$/m³"),
    "ÓLEO DIESEL": ("DIESEL_B_S500", "R$/L"),
    "OLEO DIESEL": ("DIESEL_B_S500", "R$/L"),
    "ÓLEO DIESEL S10": ("DIESEL_B_S10", "R$/L"),
    "OLEO DIESEL S10": ("DIESEL_B_S10", "R$/L"),
    "OLEO DIESEL S500": ("DIESEL_B_S500", "R$/L"),
    "DIESEL S10": ("DIESEL_B_S10", "R$/L"),
    "DIESEL B S10": ("DIESEL_B_S10", "R$/L"),
    "GLP": ("GLP", "R$/13kg"),
    "GLP P13": ("GLP", "R$/13kg"),
}
UF_NAMES = {
    "ACRE":"AC","ALAGOAS":"AL","AMAPA":"AP","AMAZONAS":"AM","BAHIA":"BA","CEARA":"CE",
    "DISTRITO FEDERAL":"DF","ESPIRITO SANTO":"ES","GOIAS":"GO","MARANHAO":"MA","MATO GROSSO":"MT",
    "MATO GROSSO DO SUL":"MS","MINAS GERAIS":"MG","PARA":"PA","PARAIBA":"PB","PARANA":"PR",
    "PERNAMBUCO":"PE","PIAUI":"PI","RIO DE JANEIRO":"RJ","RIO GRANDE DO NORTE":"RN","RIO GRANDE DO SUL":"RS",
    "RONDONIA":"RO","RORAIMA":"RR","SANTA CATARINA":"SC","SAO PAULO":"SP","SERGIPE":"SE","TOCANTINS":"TO",
}


def normalized(value: object) -> str:
    return " ".join(str(value or "").strip().upper().replace("\n", " ").split())


def unaccented(value: str) -> str:
    import unicodedata
    return "".join(char for char in unicodedata.normalize("NFD", value) if unicodedata.category(char) != "Mn")


class LinkParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)


def discover_source(html: str, base_url: str = SOURCE_INDEX) -> str:
    parser = LinkParser()
    parser.feed(html)
    candidates = [urljoin(base_url, href) for href in parser.links
                  if re.search(r"resumo_semanal_lpc_\d{4}-\d{2}-\d{2}_\d{4}-\d{2}-\d{2}\.xlsx(?:$|\?)", href, re.I)]
    unique = sorted(set(candidates))
    if not unique:
        raise ValueError("Nenhum resumo semanal XLSX reconhecível na página oficial da ANP.")
    dated = []
    for candidate in unique:
        match = re.search(r"resumo_semanal_lpc_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\.xlsx", candidate, re.I)
        if match:
            dated.append((date.fromisoformat(match.group(2)), candidate))
    if not dated:
        raise ValueError("Não foi possível identificar o período dos resumos semanais oficiais.")
    latest_end = max(item[0] for item in dated)
    latest = [url for end, url in dated if end == latest_end]
    if len(latest) != 1:
        raise ValueError(f"Mais de um resumo semanal para o período mais recente ({latest_end}).")
    return latest[0]


def request(url: str, timeout: int = 45) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": "BetoData/1.0 (+https://github.com/williamqf/Beto)"})
    with urllib.request.urlopen(req, timeout=timeout) as response:
        body = response.read()
        if response.status != 200 or not body:
            raise ValueError(f"Download inválido: HTTP {response.status}, {len(body)} bytes.")
        return body


def parse_price(value: object, maximum: float = MAX_PRICE) -> float:
    if isinstance(value, (float, int)):
        price = float(value)
    else:
        raw = str(value or "").strip().replace("R$", "").replace(" ", "")
        if "," in raw:
            raw = raw.replace(".", "").replace(",", ".")
        price = float(raw)
    if not MIN_PRICE <= price <= maximum:
        raise ValueError(f"Preço fora da faixa de sanidade ({price}).")
    return round(price, 4)


def parse_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None or str(value).strip() == "":
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Data inválida: {value!r}")


def column(headers: list[str], *names: str, required: bool = False) -> int | None:
    for name in names:
        alias = normalized(name)
        found = next((i for i, header in enumerate(headers) if header == alias or header.startswith(alias + " ") or header.startswith(alias + " (")), None)
        if found is not None:
            return found
    if required:
        raise ValueError(f"Coluna esperada ausente: {' / '.join(names)}")
    return None


def read_sheet(ws, scope: str) -> list[dict]:
    rows = ws.iter_rows(values_only=True)
    materialized = list(rows)
    header_index = next((index for index, row in enumerate(materialized[:20])
                         if "PRODUTO" in [normalized(c) for c in row]
                         and any(normalized(c).startswith("PREÇO MÉDIO") for c in row)), None)
    if header_index is None:
        return []
    headers = [normalized(c) for c in materialized[header_index]]
    fuel_i = column(headers, "PRODUTO", "COMBUSTÍVEL", required=True)
    avg_i = column(headers, "PREÇO MÉDIO REVENDA", "PREÇO MÉDIO", "PREÇO MÉDIO AO CONSUMIDOR", required=True)
    unit_i = column(headers, "UNIDADE DE MEDIDA", "UNIDADE", required=True)
    uf_i = column(headers, "UF", "ESTADO", "ESTADOS")
    city_i = column(headers, "MUNICÍPIO", "MUNICIPIO")
    min_i = column(headers, "PREÇO MÍNIMO REVENDA", "PREÇO MÍNIMO")
    max_i = column(headers, "PREÇO MÁXIMO REVENDA", "PREÇO MÁXIMO")
    sample_i = column(headers, "NÚMERO DE POSTOS PESQUISADOS", "NÚMERO DE POSTOS", "POSTOS PESQUISADOS")
    start_i = column(headers, "DATA INICIAL", "DATA INÍCIO")
    end_i = column(headers, "DATA FINAL", "DATA FIM")
    output: list[dict] = []
    for row in materialized[header_index + 1:]:
        if not row or len(row) <= max(fuel_i, avg_i, unit_i):
            continue
        fuel_label = normalized(row[fuel_i])
        if fuel_label not in FUEL_ALIASES:
            # Unknown fuels in a changed/expanded official schema are not published silently.
            if fuel_label:
                raise ValueError(f"Combustível não mapeado na fonte: {fuel_label}")
            continue
        fuel, expected_unit = FUEL_ALIASES[fuel_label]
        unit_raw = normalized(row[unit_i]).replace("³", "3").replace("MÂ³", "M3")
        unit = "R$/m³" if "M3" in unit_raw else "R$/13kg" if "13KG" in unit_raw or "13 KG" in unit_raw else "R$/L" if unit_raw in {"L", "LITRO", "R$/L"} else ""
        if unit != expected_unit:
            raise ValueError(f"Unidade inesperada para {fuel_label}: {row[unit_i]!r}")
        uf_raw = normalized(row[uf_i]) if uf_i is not None and uf_i < len(row) else ""
        uf = UF_NAMES.get(unaccented(uf_raw), uf_raw if re.fullmatch(r"[A-Z]{2}", uf_raw) else None)
        city = str(row[city_i]).strip() if city_i is not None and city_i < len(row) and row[city_i] else None
        if scope == "MUNICIPALITY" and (not city or not uf):
            raise ValueError("Município sem UF ou nome na planilha da ANP.")
        if scope == "STATE" and not uf:
            raise ValueError("Agregação estadual sem UF.")
        max_for_unit = MAX_PRICE_PER_13KG if expected_unit == "R$/13kg" else MAX_PRICE
        avg = parse_price(row[avg_i], max_for_unit)
        def optional_price(idx):
            return parse_price(row[idx], max_for_unit) if idx is not None and idx < len(row) and row[idx] not in (None, "") else None
        sample = row[sample_i] if sample_i is not None and sample_i < len(row) else None
        output.append({
            "scope": scope, "uf": uf if scope != "NATIONAL" else None,
            "municipality": city if scope == "MUNICIPALITY" else None,
            "fuelType": fuel, "unit": expected_unit, "averagePrice": avg,
            "minimumPrice": optional_price(min_i), "maximumPrice": optional_price(max_i),
            "sampleSize": int(sample) if isinstance(sample, (int, float)) and sample >= 0 else None,
            "referenceStart": parse_date(row[start_i]).isoformat() if start_i is not None and start_i < len(row) and row[start_i] else None,
            "referenceEnd": parse_date(row[end_i]).isoformat() if end_i is not None and end_i < len(row) and row[end_i] else None,
        })
    return output


def parse_workbook(raw: bytes, source_url: str) -> tuple[str, list[dict]]:
    if not raw:
        raise ValueError("Arquivo ANP vazio.")
    match = re.search(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\.xlsx", source_url, re.I)
    if not match:
        raise ValueError("Período de referência não identificado no nome oficial do arquivo.")
    start, end = date.fromisoformat(match.group(1)), date.fromisoformat(match.group(2))
    if start > end or end > date.today():
        raise ValueError("Período de referência inválido.")
    import io
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    records = []
    found = set()
    names = {normalized(s): s for s in wb.sheetnames}
    sheet_aliases = {"NATIONAL": ("BRASIL", "NACIONAL"), "STATE": ("ESTADOS", "UF"), "MUNICIPALITY": ("MUNICÍPIOS", "MUNICIPIOS")}
    for scope, aliases in sheet_aliases.items():
        sheet = next((names[normalized(a)] for a in aliases if normalized(a) in names), None)
        if sheet:
            found.add(scope)
            records.extend(read_sheet(wb[sheet], scope))
    if found != set(sheet_aliases):
        raise ValueError(f"Abas agregadas incompletas: {sorted(found)}")
    fuels = {r["fuelType"] for r in records}
    if not {"GASOLINE_C", "ETHANOL_HYDRATED"}.issubset(fuels):
        raise ValueError("Gasolina e etanol precisam existir no resumo.")
    if not any(r["scope"] == "MUNICIPALITY" for r in records):
        raise ValueError("Nenhum registro municipal encontrado.")
    if not any(r["scope"] == "STATE" and r["uf"] for r in records):
        raise ValueError("Nenhuma UF encontrada.")
    if len(records) < MIN_RECORDS:
        raise ValueError(f"Dataset pequeno demais: {len(records)} registros (mínimo {MIN_RECORDS}).")
    keys = [(r["scope"], r["uf"], r["municipality"], r["fuelType"]) for r in records]
    if len(keys) != len(set(keys)):
        raise ValueError("Duplicidade inesperada na chave escopo/local/combustível.")
    for r in records:
        if r["referenceStart"] and r["referenceStart"] != start.isoformat():
            raise ValueError("Data inicial da planilha não corresponde ao período do arquivo.")
        if r["referenceEnd"] and r["referenceEnd"] != end.isoformat():
            raise ValueError("Data final da planilha não corresponde ao período do arquivo.")
        r["referenceStart"] = start.isoformat()
        r["referenceEnd"] = end.isoformat()
    records.sort(key=lambda r: (r["scope"], r["uf"] or "", r["municipality"] or "", r["fuelType"]))
    return f"{start.isoformat()}_{end.isoformat()}", records


def build_files(version: str, records: list[dict], generated_at: str) -> dict[str, bytes]:
    grouped: dict[str, list[dict]] = {"BR": []}
    for source_record in records:
        r = {**source_record, "sourceName": SOURCE_NAME, "sourceReference": SOURCE_INDEX,
             "datasetVersion": version, "generatedAt": generated_at}
        if r["scope"] == "NATIONAL":
            grouped["BR"].append(r)
        elif r["uf"]:
            grouped.setdefault(r["uf"], []).append(r)
    files: dict[str, bytes] = {}
    entries = {}
    for region, rows in sorted(grouped.items()):
        if not rows:
            continue
        name = f"datasets/fuel/{version}/{region}.json"
        body = json.dumps({"schemaVersion": SCHEMA_VERSION, "datasetVersion": version, "region": region,
                           "records": rows}, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
        files[name] = body
        entries[region] = {"path": name, "sha256": hashlib.sha256(body).hexdigest(), "recordCount": len(rows)}
    if "BR" not in entries:
        raise ValueError("Agregação nacional ausente.")
    manifest = {"schemaVersion": SCHEMA_VERSION, "datasetVersion": version, "generatedAt": generated_at,
                "source": "ANP", "sourceName": SOURCE_NAME, "sourceIndex": SOURCE_INDEX,
                "sourceReference": SOURCE_INDEX, "referenceStart": version.split("_")[0],
                "referenceEnd": version.split("_")[1], "datasets": entries}
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2).encode()
    return files


def sanity_against_previous(files: dict[str, bytes], previous: Path) -> None:
    old_manifest_path = previous / "manifest.json"
    if not old_manifest_path.exists():
        return
    old_manifest = json.loads(old_manifest_path.read_text(encoding="utf-8"))
    new_manifest = json.loads(files["manifest.json"])
    old_count = sum(int(v.get("recordCount", 0)) for v in old_manifest.get("datasets", {}).values())
    new_count = sum(int(v.get("recordCount", 0)) for v in new_manifest.get("datasets", {}).values())
    if old_count and new_count < old_count * 0.5:
        raise ValueError(f"Queda anormal no dataset: {old_count} → {new_count} registros.")


def publish_atomically(files: dict[str, bytes], destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="beto-data-") as td:
        stage = Path(td) / "stage"
        for relative, body in files.items():
            target = stage / relative
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(body)
        backup = Path(td) / "old"
        try:
            if destination.exists():
                os.replace(destination, backup)
            os.replace(stage, destination)
        except Exception:
            if backup.exists() and not destination.exists():
                os.replace(backup, destination)
            raise


def update(fetch=request, destination: Path = DIST, generated_at: str | None = None) -> bool:
    html = fetch(SOURCE_INDEX).decode("utf-8", errors="replace")
    source_url = discover_source(html)
    raw = fetch(source_url)
    version, records = parse_workbook(raw, source_url)
    files = build_files(version, records, generated_at or datetime.now(timezone.utc).isoformat(timespec="seconds"))
    sanity_against_previous(files, destination)
    old_manifest = destination / "manifest.json"
    if old_manifest.exists():
        old = json.loads(old_manifest.read_text(encoding="utf-8"))
        new = json.loads(files["manifest.json"])
        if old.get("datasetVersion") == version and all(
                (destination / meta["path"]).exists() and hashlib.sha256((destination / meta["path"]).read_bytes()).hexdigest() == meta["sha256"]
                for meta in new["datasets"].values()):
            return False
    publish_atomically(files, destination)
    return True


def seed_previous_publication(destination: Path, fetch=request) -> None:
    """Restore the last Pages dataset into the runner so no-op and size checks work across runs."""
    if (destination / "manifest.json").exists():
        return
    base = os.environ.get("BETO_PREVIOUS_BASE_URL", "https://williamqf.github.io/beto-data").rstrip("/")
    try:
        manifest_bytes = fetch(f"{base}/manifest.json")
        manifest = json.loads(manifest_bytes)
        if manifest.get("schemaVersion") != SCHEMA_VERSION:
            return
        files = {"manifest.json": manifest_bytes}
        for entry in manifest.get("datasets", {}).values():
            path = entry.get("path", "")
            if not path.startswith("datasets/fuel/") or ".." in path:
                raise ValueError("Caminho inválido no manifesto publicado anteriormente.")
            body = fetch(f"{base}/{path}")
            if hashlib.sha256(body).hexdigest() != entry.get("sha256"):
                raise ValueError("Checksum da publicação anterior inválido.")
            files[path] = body
        publish_atomically(files, destination)
    except Exception as exc:
        # Before the first Pages deploy a 404 is normal. Upstream update still must be validated.
        print(f"Publicação anterior não recuperada (primeira execução ou fonte indisponível): {exc}", file=sys.stderr)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dist", type=Path, default=DIST)
    args = parser.parse_args()
    try:
        seed_previous_publication(args.dist)
        changed = update(destination=args.dist)
        print(json.dumps({"changed": changed, "dist": str(args.dist)}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(f"Beto Data ANP: atualização cancelada; dataset anterior preservado. {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
