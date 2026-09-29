import hashlib
import io
import json
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src import update_anp as anp


def workbook_bytes():
    wb = Workbook()
    ws = wb.active
    ws.title = "BRASIL"
    headers = ["PRODUTO", "UNIDADE DE MEDIDA", "PREÇO MÉDIO REVENDA", "DATA INICIAL", "DATA FINAL"]
    ws.append(headers)
    fuels = [("GASOLINA C", "R$/L", 6.1), ("ETANOL HIDRATADO", "R$/L", 4.1), ("GNV", "R$/m³", 4.5)]
    start, end = date(2026, 9, 20), date(2026, 9, 26)
    for name, unit, price in fuels:
        ws.append([name, unit, price, start, end])
    states = wb.create_sheet("ESTADOS")
    states.append(["PRODUTO", "UNIDADE DE MEDIDA", "PREÇO MÉDIO REVENDA", "ESTADO"])
    for uf in ("MG", "SP", "RJ"):
        for name, unit, price in fuels:
            states.append([name, unit, price, uf])
    cities = wb.create_sheet("MUNICÍPIOS")
    cities.append(["PRODUTO", "UNIDADE DE MEDIDA", "PREÇO MÉDIO REVENDA", "ESTADO", "MUNICÍPIO"])
    for i in range(12):
        for name, unit, price in fuels:
            cities.append([name, unit, price + (i / 100), "MG", f"Cidade {i:02}"])
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def mutate_workbook(change):
    wb = load_workbook(io.BytesIO(workbook_bytes()))
    change(wb)
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


class PipelineTests(unittest.TestCase):
    def test_source_discovery_requires_unambiguous_official_weekly_file(self):
        html = '<a href="/precos/arquivos-lpc/2026/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx">Resumo</a>' \
               '<a href="/precos/arquivos-lpc/2026/resumo_semanal_lpc_2026-09-13_2026-09-19.xlsx">Anterior</a>'
        self.assertIn("2026-09-26", anp.discover_source(html))
        with self.assertRaises(ValueError):
            anp.discover_source("<html>sem arquivo</html>")

    def test_valid_workbook_generates_deterministic_manifest_and_checksums(self):
        version, records = anp.parse_workbook(workbook_bytes(), "https://gov.br/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx")
        self.assertEqual("R$/m³", next(r for r in records if r["fuelType"] == "GNV")["unit"])
        first = anp.build_files(version, records, "2026-09-27T12:00:00+00:00")
        second = anp.build_files(version, records, "2026-09-27T12:00:00+00:00")
        self.assertEqual(first, second)
        manifest = json.loads(first["manifest.json"])
        self.assertEqual(manifest["schemaVersion"], 1)
        for meta in manifest["datasets"].values():
            self.assertEqual(hashlib.sha256(first[meta["path"]]).hexdigest(), meta["sha256"])
        self.assertIn("MG", manifest["datasets"])

    def test_empty_invalid_schema_price_and_date_fail(self):
        with self.assertRaises(ValueError):
            anp.parse_workbook(b"", "https://x/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx")
        with self.assertRaises(ValueError):
            anp.parse_price("619,00")
        with self.assertRaises(ValueError):
            anp.parse_date("ontem")
        self.assertTrue(anp.parse_workbook(workbook_bytes(), "https://x/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx")[1])

    def test_changed_schema_unknown_fuel_bad_price_date_and_missing_uf_fail(self):
        source = "https://x/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx"
        bad_schema = mutate_workbook(lambda wb: wb["MUNICÍPIOS"].cell(1, 1, "CAMPO DESCONHECIDO"))
        with self.assertRaises(ValueError):
            anp.parse_workbook(bad_schema, source)
        unknown = mutate_workbook(lambda wb: wb["BRASIL"].append(["COMBUSTÍVEL NOVO", "R$/L", 5.0, date.today(), date.today()]))
        with self.assertRaisesRegex(ValueError, "não mapeado"):
            anp.parse_workbook(unknown, source)
        bad_price = mutate_workbook(lambda wb: wb["MUNICÍPIOS"].cell(2, 3, 619))
        with self.assertRaisesRegex(ValueError, "sanidade"):
            anp.parse_workbook(bad_price, source)
        bad_date = mutate_workbook(lambda wb: wb["BRASIL"].cell(2, 4, "data errada"))
        with self.assertRaisesRegex(ValueError, "Data inválida"):
            anp.parse_workbook(bad_date, source)
        no_uf = mutate_workbook(lambda wb: setattr(wb["MUNICÍPIOS"]["D2"], "value", None))
        with self.assertRaisesRegex(ValueError, "Município sem UF"):
            anp.parse_workbook(no_uf, source)

    def test_duplicate_and_abnormally_smaller_dataset_fail(self):
        source = "https://x/resumo_semanal_lpc_2026-09-20_2026-09-26.xlsx"
        duplicate = mutate_workbook(lambda wb: wb["MUNICÍPIOS"].append(list(wb["MUNICÍPIOS"].values)[1]))
        with self.assertRaisesRegex(ValueError, "Duplicidade"):
            anp.parse_workbook(duplicate, source)
        version, records = anp.parse_workbook(workbook_bytes(), source)
        files = anp.build_files(version, records, "2026-09-27T12:00:00+00:00")
        with tempfile.TemporaryDirectory() as td:
            old = Path(td)
            (old / "manifest.json").write_text(json.dumps({"datasets": {"BR": {"recordCount": 1000}}}), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Queda anormal"):
                anp.sanity_against_previous(files, old)

    def test_failed_download_keeps_previous_published_manifest(self):
        with tempfile.TemporaryDirectory() as td:
            destination = Path(td) / "dist"
            destination.mkdir()
            (destination / "manifest.json").write_text('{"datasetVersion":"old"}', encoding="utf-8")
            def fail(_url):
                raise OSError("offline")
            with self.assertRaises(OSError):
                anp.update(fetch=fail, destination=destination)
            self.assertEqual((destination / "manifest.json").read_text(), '{"datasetVersion":"old"}')

    def test_atomic_publish_preserves_files_as_a_complete_tree(self):
        with tempfile.TemporaryDirectory() as td:
            dest = Path(td) / "dist"
            anp.publish_atomically({"manifest.json": b"{}", "datasets/fuel/test/MG.json": b"{}"}, dest)
            self.assertEqual((dest / "manifest.json").read_bytes(), b"{}")
            self.assertTrue((dest / "datasets/fuel/test/MG.json").exists())


if __name__ == "__main__":
    unittest.main()
